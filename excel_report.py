from __future__ import annotations

from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MONEY_FMT = '"$" #,##0.00'
DATE_FMT = "DD/MM/YYYY"
PCT_FMT = "0.00%"

FONT_TITLE = Font(bold=True, size=14, color="0F172A")
FONT_META = Font(size=10, color="64748B")
FONT_HEADER = Font(bold=True, size=10, color="FFFFFF")
FONT_SECTION = Font(bold=True, size=11, color="FFFFFF")
FONT_BOLD = Font(bold=True)

FILL_HEADER = PatternFill("solid", fgColor="334155")
FILL_SECTION_90 = PatternFill("solid", fgColor="7F1D1D")
FILL_SECTION_60 = PatternFill("solid", fgColor="B91C1C")
FILL_SECTION_30 = PatternFill("solid", fgColor="D97706")
FILL_TOTAL = PatternFill("solid", fgColor="E2E8F0")
FILL_SUMMARY_TOTAL = PatternFill("solid", fgColor="DCE6F1")
FILL_RED = PatternFill("solid", fgColor="FEE2E2")
FILL_AMBER = PatternFill("solid", fgColor="FEF3C7")
FILL_GREEN = PatternFill("solid", fgColor="DCFCE7")

THIN = Side(style="thin", color="000000")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_BOTTOM = Border(bottom=Side(style="thin", color="E2E8F0"))

BUCKETS = ["0 - 30", "31 - 60", "61 - 90", "91+"]


def bucket_of(dias: int | None) -> str:
    """Tramo según días desde la venta. Sin fecha de documento se considera en fecha."""
    d = dias or 0
    if d <= 30:
        return BUCKETS[0]
    if d <= 60:
        return BUCKETS[1]
    if d <= 90:
        return BUCKETS[2]
    return BUCKETS[3]


def bucket_totals(invoices: list[dict]) -> dict[str, float]:
    """{tramo: suma de saldo} más la clave "Total"."""
    totals = {b: 0.0 for b in BUCKETS}
    for it in invoices:
        totals[it["tramo"]] += it["saldo"]
    totals["Total"] = sum(totals[b] for b in BUCKETS)
    return totals


def _days_since(d: date | None, today: date) -> int | None:
    return (today - d).days if d else None


def _invoice_rows(clientes: list[dict], today: date) -> list[dict]:
    """Facturas abiertas (saldo > 0) aplanadas con datos del cliente, vendedor, días y tramo."""
    rows = []
    for c in clientes:
        for f in c.get("facturas") or []:
            saldo = float(f.get("SaldoPendiente") or 0)
            if saldo <= 0:
                continue
            doc_date = f.get("_DocDate")
            due_date = f.get("_DocDueDate")
            dias_venta = _days_since(doc_date, today)
            dias_vencida = _days_since(due_date, today)
            rows.append(
                {
                    "vendedor": c.get("vendedor") or "Sin Vendedor Asignado",
                    "cliente": c.get("CardName") or "",
                    "codigo": c.get("CardCode") or "",
                    "documento": f.get("NumAtCard") or "",
                    "fecha_doc": doc_date,
                    "vencimiento": due_date,
                    "dias_venta": dias_venta,
                    "tramo": bucket_of(dias_venta),
                    "dias_vencida": dias_vencida if dias_vencida and dias_vencida > 0 else 0,
                    "total": round(float(f.get("DocTotal") or 0), 2),
                    "pagado": round(float(f.get("PaidToDate") or 0), 2),
                    "saldo": round(saldo, 2),
                }
            )
    return rows


def _write_title(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = FONT_TITLE
    ws["A2"] = subtitle
    ws["A2"].font = FONT_META


def _write_header(ws, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _set_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_bucket_summary(
    ws,
    row: int,
    rows_data: list[tuple[str, dict[str, float]]],
    label_header: str = "Etiquetas de fila",
    include_total: bool = True,
) -> int:
    """
    Tabla de resumen por tramos (estilo tabla dinámica): una fila por etiqueta,
    "Total general" (opcional) y "Total en %". Devuelve la próxima fila libre.
    """
    headers = [label_header] + [f"Suma de {b}" for b in BUCKETS] + ["Suma de Saldo vencido"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = FONT_BOLD
        cell.border = BORDER_ALL
    row += 1

    def write_amounts(r: int, label: str, totals: dict[str, float], bold: bool, fill: PatternFill | None):
        ws.cell(row=r, column=1, value=label)
        values = [totals[b] for b in BUCKETS] + [totals["Total"]]
        for col, v in enumerate(values, start=2):
            cell = ws.cell(row=r, column=col, value=round(v, 2) if v else None)
            cell.number_format = MONEY_FMT
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = BORDER_ALL
            cell.font = FONT_BOLD
            if fill:
                cell.fill = fill

    for label, totals in rows_data:
        write_amounts(row, label, totals, bold=True, fill=None)
        row += 1

    grand = {b: sum(t[b] for _, t in rows_data) for b in BUCKETS}
    grand["Total"] = sum(grand[b] for b in BUCKETS)
    if include_total:
        write_amounts(row, "Total general", grand, bold=True, fill=FILL_SUMMARY_TOTAL)
        row += 1

    ws.cell(row=row, column=1, value="Total en %")
    pct_values = [(grand[b] / grand["Total"]) if grand["Total"] else 0 for b in BUCKETS] + [1 if grand["Total"] else 0]
    for col, v in enumerate(pct_values, start=2):
        ws.cell(row=row, column=col, value=v).number_format = PCT_FMT
    for col in range(1, len(headers) + 1):
        ws.cell(row=row, column=col).font = FONT_BOLD
        ws.cell(row=row, column=col).alignment = Alignment(horizontal="right") if col > 1 else Alignment()
    return row + 1


# ---------------------------------------------------------------- tablas de facturas

URGENT_HEADERS = ["Cliente", "Código", "Documento", "Fecha doc.", "Vencimiento", "Días desde venta", "Días vencida", "Saldo"]


def _write_urgent_section(ws, row: int, title: str, fill: PatternFill, items: list[dict]) -> int:
    ws.cell(row=row, column=1, value=f"{title} ({len(items)} facturas)")
    for col in range(1, len(URGENT_HEADERS) + 1):
        ws.cell(row=row, column=col).fill = fill
        ws.cell(row=row, column=col).font = FONT_SECTION
    row += 1
    _write_header(ws, row, URGENT_HEADERS)
    row += 1

    if not items:
        ws.cell(row=row, column=1, value="Sin facturas en este tramo.").font = FONT_META
        return row + 2

    first = row
    for it in items:
        values = [
            it["cliente"], it["codigo"], it["documento"], it["fecha_doc"], it["vencimiento"],
            it["dias_venta"], it["dias_vencida"], it["saldo"],
        ]
        for col, v in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=v).border = BORDER_BOTTOM
        ws.cell(row=row, column=4).number_format = DATE_FMT
        ws.cell(row=row, column=5).number_format = DATE_FMT
        ws.cell(row=row, column=8).number_format = MONEY_FMT
        row += 1

    ws.cell(row=row, column=1, value="Total")
    ws.cell(row=row, column=8, value=f"=SUM(H{first}:H{row - 1})").number_format = MONEY_FMT
    for col in range(1, len(URGENT_HEADERS) + 1):
        ws.cell(row=row, column=col).fill = FILL_TOTAL
        ws.cell(row=row, column=col).font = FONT_BOLD
    return row + 3


def _write_invoice_table(ws, header_row: int, invoices: list[dict], include_vendor: bool) -> None:
    """Tabla plana de facturas con filtros, encabezado fijo y totales con SUBTOTAL."""
    headers = [
        "Cliente", "Código", "Documento", "Fecha doc.", "Vencimiento", "Días desde venta", "Tramo",
        "Días vencida", "Total", "Pagado", "Saldo", "Estado",
    ]
    if include_vendor:
        headers = ["Vendedor"] + headers
    off = 1 if include_vendor else 0
    _write_header(ws, header_row, headers)

    if include_vendor:
        items = sorted(invoices, key=lambda i: (i["vendedor"], i["cliente"], -(i["dias_venta"] or 0)))
    else:
        items = sorted(invoices, key=lambda i: (i["cliente"], -(i["dias_venta"] or 0)))

    money_cols = (9 + off, 10 + off, 11 + off)
    row = header_row + 1
    for it in items:
        estado = "Vencida" if it["dias_vencida"] > 0 else "En fecha"
        values = [
            it["cliente"], it["codigo"], it["documento"], it["fecha_doc"], it["vencimiento"],
            it["dias_venta"], it["tramo"], it["dias_vencida"], it["total"], it["pagado"], it["saldo"], estado,
        ]
        if include_vendor:
            values = [it["vendedor"]] + values
        for col, v in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=v).border = BORDER_BOTTOM
        ws.cell(row=row, column=4 + off).number_format = DATE_FMT
        ws.cell(row=row, column=5 + off).number_format = DATE_FMT
        for col in money_cols:
            ws.cell(row=row, column=col).number_format = MONEY_FMT

        tramo_cell = ws.cell(row=row, column=7 + off)
        tramo_cell.alignment = Alignment(horizontal="center")
        if it["tramo"] == BUCKETS[3]:
            tramo_cell.fill = FILL_RED
            tramo_cell.font = FONT_BOLD
        elif it["tramo"] == BUCKETS[2]:
            tramo_cell.fill = FILL_RED
        elif it["tramo"] == BUCKETS[1]:
            tramo_cell.fill = FILL_AMBER
        estado_cell = ws.cell(row=row, column=12 + off)
        estado_cell.fill = FILL_RED if estado == "Vencida" else FILL_GREEN
        estado_cell.alignment = Alignment(horizontal="center")
        row += 1

    last = row - 1
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(last, header_row)}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2 + off)

    ws.cell(row=row, column=1, value="Totales")
    if last >= header_row + 1:
        for col in money_cols:
            letter = get_column_letter(col)
            ws.cell(row=row, column=col, value=f"=SUBTOTAL(9,{letter}{header_row + 1}:{letter}{last})").number_format = MONEY_FMT
    for col in range(1, len(headers) + 1):
        ws.cell(row=row, column=col).fill = FILL_TOTAL
        ws.cell(row=row, column=col).font = FONT_BOLD


# ---------------------------------------------------------------- Excel de vendedores

def _build_urgent_sheet(ws, invoices: list[dict], vendedor_name: str, generated_at: str) -> None:
    _write_title(ws, "Facturas urgentes por antigüedad de venta", f"Vendedor: {vendedor_name}   |   Generado: {generated_at}")
    _set_widths(ws, [40, 18, 18, 18, 18, 18, 18, 18])

    row = _write_bucket_summary(
        ws, 4, [(vendedor_name, bucket_totals(invoices))], label_header="Vendedor", include_total=False
    )

    def key(it):
        return (it["cliente"], -(it["dias_venta"] or 0))

    def by_bucket(b):
        return sorted((i for i in invoices if i["tramo"] == b), key=key)

    row += 2
    row = _write_urgent_section(ws, row, "1) Más de 90 días desde la venta", FILL_SECTION_90, by_bucket(BUCKETS[3]))
    row = _write_urgent_section(ws, row, "2) Entre 61 y 90 días desde la venta", FILL_SECTION_60, by_bucket(BUCKETS[2]))
    _write_urgent_section(ws, row, "3) Entre 31 y 60 días desde la venta", FILL_SECTION_30, by_bucket(BUCKETS[1]))


def _build_full_sheet(ws, invoices: list[dict], vendedor_name: str, generated_at: str) -> None:
    _write_title(ws, "Reporte completo de facturas abiertas", f"Vendedor: {vendedor_name}   |   Generado: {generated_at}")
    _set_widths(ws, [40, 10, 20, 13, 13, 12, 10, 12, 18, 18, 18, 12])
    _write_invoice_table(ws, 4, invoices, include_vendor=False)


def generate_excel_report(output_path: str, vendedor_name: str, generated_at: str, clientes: list[dict]) -> None:
    invoices = _invoice_rows(clientes, date.today())

    wb = Workbook()
    urgent = wb.active
    urgent.title = "Urgente"
    _build_urgent_sheet(urgent, invoices, vendedor_name, generated_at)

    full = wb.create_sheet("Reporte completo")
    _build_full_sheet(full, invoices, vendedor_name, generated_at)

    wb.save(output_path)


# ---------------------------------------------------------------- Excel de gerentes

def _build_manager_summary_sheet(ws, invoices: list[dict], generated_at: str) -> None:
    _write_title(ws, f"Status de Deuda al {date.today():%d-%m-%y}", f"Generado: {generated_at}")
    _set_widths(ws, [32, 20, 20, 20, 20, 24])

    por_vendedor: dict[str, list[dict]] = defaultdict(list)
    for it in invoices:
        por_vendedor[it["vendedor"]].append(it)
    rows_data = [(v, bucket_totals(items)) for v, items in sorted(por_vendedor.items(), key=lambda kv: kv[0].lower())]
    _write_bucket_summary(ws, 4, rows_data)


def _build_manager_client_sheet(ws, invoices: list[dict], generated_at: str) -> None:
    _write_title(ws, "Deuda por cliente y tramo", f"Generado: {generated_at}")
    _set_widths(ws, [26, 40, 10, 18, 18, 18, 18, 20])

    headers = ["Vendedor", "Cliente", "Código"] + BUCKETS + ["Total"]
    header_row = 4
    _write_header(ws, header_row, headers)

    por_cliente: dict[tuple[str, str, str], list[dict]] = defaultdict(list)
    for it in invoices:
        por_cliente[(it["vendedor"], it["cliente"], it["codigo"])].append(it)
    data = [(k, bucket_totals(items)) for k, items in por_cliente.items()]
    data.sort(key=lambda kt: (kt[0][0].lower(), -kt[1]["Total"]))

    row = header_row + 1
    for (vendedor, cliente, codigo), totals in data:
        values = [vendedor, cliente, codigo] + [totals[b] or None for b in BUCKETS] + [totals["Total"]]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = BORDER_BOTTOM
            if col >= 4:
                cell.number_format = MONEY_FMT
        if totals[BUCKETS[3]] or totals[BUCKETS[2]]:
            ws.cell(row=row, column=2).font = Font(bold=True, color="B91C1C")
        row += 1

    last = row - 1
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(last, header_row)}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)

    ws.cell(row=row, column=1, value="Totales")
    if last >= header_row + 1:
        for col in range(4, len(headers) + 1):
            letter = get_column_letter(col)
            ws.cell(row=row, column=col, value=f"=SUBTOTAL(9,{letter}{header_row + 1}:{letter}{last})").number_format = MONEY_FMT
    for col in range(1, len(headers) + 1):
        ws.cell(row=row, column=col).fill = FILL_TOTAL
        ws.cell(row=row, column=col).font = FONT_BOLD


def generate_manager_excel(output_path: str, generated_at: str, clientes: list[dict]) -> None:
    invoices = _invoice_rows(clientes, date.today())

    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"
    _build_manager_summary_sheet(resumen, invoices, generated_at)

    _build_manager_client_sheet(wb.create_sheet("Por cliente"), invoices, generated_at)

    detalle = wb.create_sheet("Detalle")
    _write_title(detalle, "Detalle de facturas abiertas", f"Generado: {generated_at}")
    _set_widths(detalle, [26, 40, 10, 20, 13, 13, 12, 10, 12, 18, 18, 18, 12])
    _write_invoice_table(detalle, 4, invoices, include_vendor=True)

    wb.save(output_path)
